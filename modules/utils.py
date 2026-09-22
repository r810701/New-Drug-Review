# -*- coding: utf-8 -*-
"""
modules/utils.py
=================
共用小工具：檔案指紋（供快取失效判斷）、PubMed/DOI 連結組裝、
文字清理、JSON 安全解析等。不依賴 Streamlit，方便單元測試。
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any, Optional

# ---------------------------------------------------------------------------
# 檔案指紋：用來偵測「管理藥師是否更新過知識庫檔案」
# ---------------------------------------------------------------------------

def compute_dir_fingerprint(dir_path: Path, patterns: tuple[str, ...] = ("*",)) -> str:
    if not dir_path.exists():
        return "EMPTY"
    entries = []
    for pattern in patterns:
        for p in sorted(dir_path.rglob(pattern)):
            if p.is_file():
                stat = p.stat()
                entries.append(f"{p.relative_to(dir_path)}|{stat.st_mtime_ns}|{stat.st_size}")
    joined = "\n".join(entries)
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()[:16]

# ---------------------------------------------------------------------------
# PubMed / DOI 連結組裝
# ---------------------------------------------------------------------------

def pubmed_url(pmid: str) -> str:
    pmid = pmid.strip()
    return f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"

def doi_url(doi: str) -> str:
    doi = doi.strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi, flags=re.IGNORECASE)
    return f"https://doi.org/{doi}"

def autolink_citation(raw: str) -> str:
    raw = raw.strip()
    m = re.search(r"PMID[:\s]*([0-9]{4,9})", raw, flags=re.IGNORECASE)
    if m:
        return f"{raw} → {pubmed_url(m.group(1))}"
    m = re.search(r"DOI[:\s]*(10\.\S+)", raw, flags=re.IGNORECASE)
    if m:
        return f"{raw} → {doi_url(m.group(1))}"
    return raw

# ---------------------------------------------------------------------------
# 文字/JSON 處理
# ---------------------------------------------------------------------------

def strip_code_fences(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(json)?", "", text, flags=re.IGNORECASE).strip()
    text = re.sub(r"```$", "", text).strip()
    return text

def safe_json_loads(text: str) -> Optional[Any]:
    cleaned = strip_code_fences(text)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        m = re.search(r"(\{.*\}|\[.*\])", cleaned, flags=re.DOTALL)
        if m:
            try:
                return json.loads(m.group(1))
            except json.JSONDecodeError:
                return None
        return None

def abbreviate_common_terms(text: str) -> str:
    mapping = {
        "JAK inhibitor": "JAKi",
        "SGLT2 inhibitor": "SGLT2i",
        "Rheumatoid Arthritis": "RA",
        "Ulcerative Colitis": "UC",
        "Crohn's disease": "CD",
        "Psoriatic Arthritis": "PsA",
        "Ankylosing Spondylitis": "AS",
    }
    for full, abbr in mapping.items():
        text = re.sub(re.escape(full), abbr, text, flags=re.IGNORECASE)
    return text

def truncate(text: str, n: int = 60) -> str:
    text = (text or "").replace("\n", " ")
    return text if len(text) <= n else text[: n - 1] + "…"

# ---------------------------------------------------------------------------
# 使用者上傳檔案的文字擷取
# ---------------------------------------------------------------------------
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
TEXT_EXTRACTABLE_EXTS = {".pdf", ".txt", ".md", ".docx", ".xlsx", ".csv"}

def _extract_pdf_text_full(path: Path) -> str:
    try:
        import pypdf
    except ImportError:
        return "（系統尚未安裝 pypdf，無法解析此 PDF，請執行 `pip install pypdf`）"
    try:
        reader = pypdf.PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:  # noqa: BLE001
        return f"（PDF 解析失敗：{e}；此檔案內容 AI 無法讀取，請人工確認）"

def _extract_docx_text(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return "（系統尚未安裝 python-docx，無法解析 Word 檔，請執行 `pip install python-docx`）"
    try:
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells_text = " | ".join(c.text.strip() for c in row.cells)
                if cells_text.strip(" |"):
                    parts.append(cells_text)
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        return f"（Word 檔解析失敗：{e}；此檔案內容 AI 無法讀取，請人工確認）"

def _extract_excel_text(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        return "（不支援舊版 .xls 格式，請於 Excel 另存新檔為 .xlsx 後重新上傳）"
    try:
        import openpyxl
    except ImportError:
        return "（系統尚未安裝 openpyxl，無法解析 Excel 檔）"
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[工作表：{ws.title}]")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if any(v is not None and str(v).strip() != "" for v in row):
                    lines.append(" | ".join("" if v is None else str(v) for v in row))
                    row_count += 1
                if row_count > 300:
                    lines.append("...(此工作表列數過多，已截斷)")
                    break
        return "\n".join(lines)
    except Exception as e:  # noqa: BLE001
        return f"（Excel 檔解析失敗：{e}；此檔案內容 AI 無法讀取，請人工確認）"

def extract_relevant_snippets(
    full_text: str, keywords: list[str], max_chars: int = 8000, context_chars: int = 500,
) -> Optional[str]:
    if not full_text or not keywords:
        return None
    lower_text = full_text.lower()
    hit_spans: list[tuple[int, int]] = []
    for kw in keywords:
        kw_lower = kw.lower()
        start = 0
        while True:
            idx = lower_text.find(kw_lower, start)
            if idx == -1:
                break
            hit_spans.append((max(0, idx - context_chars), min(len(full_text), idx + len(kw) + context_chars)))
            start = idx + len(kw)
    if not hit_spans:
        return None
    hit_spans.sort()
    merged: list[list[int]] = []
    for s, e in hit_spans:
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    parts, total = [], 0
    for s, e in merged:
        snippet = full_text[s:e].strip()
        if not snippet:
            continue
        if total + len(snippet) > max_chars:
            remain = max_chars - total
            if remain <= 0:
                break
            snippet = snippet[:remain]
        parts.append(snippet)
        total += len(snippet)
        if total >= max_chars:
            break
    return "\n...\n".join(parts)

def extract_text_from_upload(
    path: Path, max_chars: int = 8000, keywords: Optional[list[str]] = None,
) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf_text_full(path)
    elif suffix in (".txt", ".md"):
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".docx":
        text = _extract_docx_text(path)
    elif suffix in (".xlsx", ".xls"):
        text = _extract_excel_text(path)
    elif suffix == ".csv":
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif suffix in IMAGE_EXTS:
        return "（圖片檔案：將以視覺方式直接提供給 AI 辨識，非文字擷取）"
    else:
        return f"（不支援自動擷取 .{suffix.lstrip('.')} 檔案的文字內容，AI 無法讀取此檔案）"

    if text.startswith("（") and text.endswith("）"):
        return text

    text = text.strip()
    if not text:
        return "（此檔案擷取不到文字，可能是掃描影像型 PDF 或空白文件；AI 無法讀取實際內容，" \
               "請人工確認或手動輸入摘要，否則 AI 只會憑一般知識作答，可能與本篇文獻不符）"

    if keywords:
        smart = extract_relevant_snippets(text, keywords, max_chars=max_chars)
        if smart:
            return smart + f"\n...(已依關鍵字智慧擷取相關段落，原文共 {len(text)} 字元)"

    if len(text) > max_chars:
        text = text[:max_chars] + f"\n...(內容過長已截斷，原文共 {len(text)} 字元，AI 僅看得到前 {max_chars} 字元)"
    return text

def excel_to_csv_text(file_obj) -> str:
    """
    把上傳的 Excel 檔案（Streamlit UploadedFile 或檔案路徑）轉成 CSV 文字，
    格式跟 Google 試算表匯出的 CSV 完全相容，可以直接沿用
    structured_data.parse_csv_text() / parse_csv_two_row_header() 既有邏輯，
    不用為了「上傳Excel」這個來源另外寫一套解析規則。
    只讀取第一個工作表；儲存格值一律轉成字串，空白格轉成空字串。
    """
    import openpyxl
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.worksheets[0]
    f = io.StringIO()
    writer = csv.writer(f)
    for row in ws.iter_rows(values_only=True):
        writer.writerow(["" if v is None else str(v) for v in row])
    return f.getvalue()

# ---------------------------------------------------------------------------
# 圖片檔案：以視覺（vision）方式直接交給 AI 辨識
# ---------------------------------------------------------------------------
_MEDIA_TYPE_BY_SUFFIX = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
}

def encode_image_for_claude(path: Path, max_dimension: int = 1568) -> tuple[str, str]:
    import base64
    try:
        from PIL import Image
        import io as _io
        img = Image.open(path)
        img = img.convert("RGB")
        w, h = img.size
        scale = min(1.0, max_dimension / max(w, h))
        if scale < 1.0:
            img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))))
        buf = _io.BytesIO()
        img.save(buf, format="JPEG", quality=85)
        return base64.b64encode(buf.getvalue()).decode("ascii"), "image/jpeg"
    except ImportError:
        data = path.read_bytes()
        media_type = _MEDIA_TYPE_BY_SUFFIX.get(path.suffix.lower(), "image/jpeg")
        return base64.b64encode(data).decode("ascii"), media_type

def load_image_for_gemini(path: Path, max_dimension: int = 1568):
    try:
        from PIL import Image
        img = Image.open(path).convert("RGB")
        w, h = img.size
        scale = min(1.0, max_dimension / max(w, h))
        if scale < 1.0:
            img = img.resize((max(1, int(w * scale)), max(1, int(h * scale))))
        return img
    except ImportError:
        return None

# ---------------------------------------------------------------------------
# 使用者貼的連結：真正抓取內容
# ---------------------------------------------------------------------------
_URL_RE = re.compile(r"https?://[^\s\u3000]+")
_GSHEET_RE = re.compile(r"https://docs\.google\.com/spreadsheets/d/([a-zA-Z0-9_-]+)")

def _to_gsheet_csv_url(url: str) -> Optional[str]:
    m = _GSHEET_RE.search(url)
    if not m:
        return None
    sheet_id = m.group(1)
    gid_match = re.search(r"[?&#]gid=(\d+)", url)
    gid = gid_match.group(1) if gid_match else "0"
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"

def _strip_html(html: str) -> str:
    html = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    html = re.sub(r"(?s)<[^>]+>", "\n", html)
    html = re.sub(r"[ \t]+", " ", html)
    html = re.sub(r"\n{2,}", "\n", html)
    return html.strip()

def fetch_url_content(url: str, max_chars: int = 6000, timeout: int = 10) -> str:
    try:
        import requests
    except ImportError:
        return "（系統尚未安裝 requests，無法抓取連結內容，請執行 `pip install requests`）"

    csv_url = _to_gsheet_csv_url(url)
    fetch_url = csv_url or url
    try:
        resp = requests.get(fetch_url, timeout=timeout, headers={"User-Agent": "Mozilla/5.0"})
    except requests.exceptions.RequestException as e:
        return f"（連結抓取失敗：{e}；請確認網址可公開存取，或改用檔案上傳）"

    if resp.status_code != 200:
        hint = "（Google 試算表需設定為「知道連結的人皆可檢視」才能被系統讀取）" if csv_url else ""
        return f"（連結回應狀態碼 {resp.status_code}，無法取得內容{hint}；請確認權限或改用檔案上傳）"

    content_type = resp.headers.get("Content-Type", "")
    if csv_url or "csv" in content_type or "text/plain" in content_type:
        text = resp.text
    elif "html" in content_type:
        text = _strip_html(resp.text)
    else:
        text = resp.text

    text = text.strip()
    if not text:
        return "（連結內容為空，或該頁面需要登入才能檢視；AI 無法讀取，請改用檔案上傳）"
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n...(內容過長已截斷，原文共 {len(text)} 字元)"
    return text

def extract_urls(text: str) -> list[str]:
    return _URL_RE.findall(text or "")

def strip_urls(text: str) -> str:
    return _URL_RE.sub("", text or "").strip()
